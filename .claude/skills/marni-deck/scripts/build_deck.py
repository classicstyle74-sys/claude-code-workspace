#!/usr/bin/env python3
"""Marni テンプレート PPTX 生成スクリプト。

テンプレート (assets/Marni_Template.pptx) のスライドを複製し、JSON スペックの
内容を流し込む。テンプレート由来の 10 タイプに加え、marni_kit のデザイン部品
(チャート・KPI タイル・プロセス図・比較・タイムライン・表・全面画像) による
拡張タイプを持つ。デザイントークン (色/フォント/ロゴ/フッター) は不変。

使い方:
    pip install python-pptx   # 未導入の場合
    python3 build_deck.py spec.json [-o output.pptx]

スペック形式は references/slide-types.md を参照。
"""

import argparse
import json
import sys
import unicodedata
from pathlib import Path

from pptx import Presentation

import marni_kit as kit

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "assets" / "Marni_Template.pptx"

# テンプレート内のスライド番号 (0 始まり) とシェイプ名の対応表。
# シェイプ名は Google Slides エクスポート由来で固定。
SLIDE_TYPES = {
    "cover": {
        "index": 0,
        "shapes": {"title": "Google Shape;133;p18"},
    },
    "toc": {
        "index": 1,
        "shapes": {
            "header": "Google Shape;140;p19",
            "secnum": "Google Shape;141;p19",
            "col_num": "Google Shape;142;p19",
            "col_title": "Google Shape;143;p19",
            "col_page": "Google Shape;144;p19",
        },
    },
    "section_red": {
        "index": 2,
        "shapes": {
            "header": "Google Shape;149;p20",
            "secnum": "Google Shape;151;p20",
        },
    },
    "section_black": {
        "index": 3,
        "shapes": {
            "header": "Google Shape;157;p21",
            "secnum": "Google Shape;159;p21",
        },
    },
    "text_image": {
        "index": 4,
        "shapes": {
            "header": "Google Shape;165;p22",
            "body": "Google Shape;166;p22",
            "image": "Google Shape;167;p22",
            "secnum": "Google Shape;168;p22",
        },
    },
    "two_images": {
        "index": 5,
        "shapes": {
            "header": "Google Shape;177;p23",
            "caption": "Google Shape;175;p23",
            "image_left": "Google Shape;174;p23",
            "image_right": "Google Shape;176;p23",
            "secnum": "Google Shape;178;p23",
        },
    },
    "three_columns": {
        "index": 6,
        "shapes": {
            "header": "Google Shape;184;p24",
            "band_title": "Google Shape;185;p24",
            "subtitles": [
                "Google Shape;186;p24",
                "Google Shape;187;p24",
                "Google Shape;188;p24",
            ],
            "images": [
                "Google Shape;189;p24",
                "Google Shape;190;p24",
                "Google Shape;191;p24",
            ],
            "bodies": [
                "Google Shape;192;p24",
                "Google Shape;193;p24",
                "Google Shape;194;p24",
            ],
            "secnum": "Google Shape;195;p24",
        },
    },
    "quote": {
        "index": 7,
        "shapes": {
            "label": "Google Shape;203;p25",
            "quote": "Google Shape;202;p25",
        },
    },
    "palette": {
        "index": 8,
        "shapes": {
            "header": "Google Shape;209;p26",
            "secnum": "Google Shape;210;p26",
        },
    },
    "closing": {
        "index": 9,
        "shapes": {"title": "Google Shape;229;p27"},
    },
}

# marni_kit で組み立てる拡張タイプ
KIT_TYPES = (
    "bullets",
    "chart",
    "stats",
    "flow",
    "comparison",
    "timeline",
    "table",
    "full_image",
)


def display_width(text):
    """全角 2 / 半角 1 の表示幅。文字数目安は半角換算のため。"""
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in text)


def warn_if_long(warnings, label, text, limit):
    if not text:
        return
    width = display_width(text)
    if width > limit:
        warnings.append(
            f"{label} が半角換算 {width} 文字相当あり目安 ({limit} 文字) を超過。"
            "はみ出す可能性があるので内容の分割を検討。"
        )


def insert_image(shape, image_path, warnings):
    if not image_path:
        return
    path = Path(image_path)
    if not path.is_file():
        warnings.append(f"画像が見つからないため空欄のままにした: {image_path}")
        return
    try:
        shape.insert_picture(str(path))
    except AttributeError:
        warnings.append(f"シェイプ {shape.name!r} は画像プレースホルダーではない")


def build_template_slide(slide, spec, warnings):
    """テンプレート由来タイプ: プレースホルダー置換で組み立てる。"""
    stype = spec["type"]
    names = SLIDE_TYPES[stype]["shapes"]

    def put(key, value, limit=None, label=None):
        if value is None:
            return
        if limit:
            warn_if_long(warnings, label or f"{stype}.{key}", value, limit)
        kit.fill_text(kit.find_shape(slide, names[key]), value)

    if stype in ("cover", "closing"):
        put("title", spec.get("title"), 60)

    elif stype == "toc":
        put("header", spec.get("title", "Table of Content"), 50)
        put("secnum", spec.get("secnum", "0.0"))
        items = spec.get("items", [])
        if len(items) > 9:
            warnings.append(f"目次が {len(items)} 行。9 行を超えるとはみ出す。")
        for it in items:
            warn_if_long(warnings, "toc 項目", it.get("title", ""), 45)
        put("col_num", "\n".join(it.get("num", "") for it in items))
        put("col_title", "\n".join(it.get("title", "") for it in items))
        put("col_page", "\n".join(it.get("page", "") for it in items))

    elif stype in ("section_red", "section_black", "palette"):
        put("header", spec.get("title"), 50)
        put("secnum", spec.get("secnum"))

    elif stype == "text_image":
        put("header", spec.get("title"), 50)
        put("secnum", spec.get("secnum"))
        put("body", spec.get("body"), 700)
        insert_image(kit.find_shape(slide, names["image"]), spec.get("image"), warnings)

    elif stype == "two_images":
        put("header", spec.get("title"), 50)
        put("secnum", spec.get("secnum"))
        put("caption", spec.get("caption"), 40)
        for key, img in zip(("image_left", "image_right"), spec.get("images", [])):
            insert_image(kit.find_shape(slide, names[key]), img, warnings)

    elif stype == "three_columns":
        put("header", spec.get("title"), 50)
        put("secnum", spec.get("secnum"))
        put("band_title", spec.get("band_title"), 60)
        columns = spec.get("columns", [])
        if len(columns) > 3:
            warnings.append("three_columns の列は最大 3。4 列目以降は無視。")
        for i, col in enumerate(columns[:3]):
            if col.get("subtitle") is not None:
                warn_if_long(warnings, f"列{i + 1} subtitle", col["subtitle"], 45)
                kit.fill_text(
                    kit.find_shape(slide, names["subtitles"][i]), col["subtitle"]
                )
            if col.get("body") is not None:
                warn_if_long(warnings, f"列{i + 1} body", col["body"], 120)
                kit.fill_text(kit.find_shape(slide, names["bodies"][i]), col["body"])
            insert_image(
                kit.find_shape(slide, names["images"][i]), col.get("image"), warnings
            )

    elif stype == "quote":
        put("label", spec.get("label", ""))
        put("quote", spec.get("quote"), 220)


def build_kit_slide(prs, template_slides, spec, warnings):
    """拡張タイプ: キャンバス + marni_kit 部品で組み立てる。"""
    stype = spec["type"]

    if stype == "full_image":
        slide = kit.canvas(prs, template_slides, title=None)
        path = Path(spec.get("image", ""))
        if path.is_file():
            kit.add_full_image(slide, path, spec.get("caption"))
        else:
            warnings.append(f"full_image の画像が見つからない: {spec.get('image')}")
        return slide

    title = spec.get("title", "")
    warn_if_long(warnings, f"{stype}.title", title, 50)
    slide = kit.canvas(prs, template_slides, secnum=spec.get("secnum", ""), title=title)
    y = kit.CONTENT_Y

    lead = spec.get("lead")
    if lead:
        warn_if_long(warnings, f"{stype}.lead", lead, 170)
        kit.add_text(slide, kit.CONTENT_X, y, kit.CONTENT_W, 0.5, lead, size=11)
        y += 0.6

    if stype == "bullets":
        items = spec.get("items", [])
        if len(items) > 8:
            warnings.append(f"bullets が {len(items)} 項目。8 項目までが目安。")
        for it in items:
            if isinstance(it, dict):
                warn_if_long(warnings, "bullets.head", it.get("head", ""), 40)
                warn_if_long(warnings, "bullets.text", it.get("text", ""), 170)
        kit.add_bullets(
            slide, items, y=y, h=kit.CONTENT_H - (y - kit.CONTENT_Y),
            columns=spec.get("columns"),
        )

    elif stype == "chart":
        body = spec.get("body")
        h = kit.CONTENT_H - (y - kit.CONTENT_Y) - (0.3 if spec.get("source") else 0)
        if body:
            warn_if_long(warnings, "chart.body", body, 380)
            kit.add_text(slide, kit.CONTENT_X, y, 2.2, h, body, size=10)
            kit.add_chart(
                slide, spec.get("chart_type", "column"), spec["categories"],
                spec["series"], x=kit.CONTENT_X + 2.5, y=y,
                w=kit.CONTENT_W - 2.5, h=h,
            )
        else:
            kit.add_chart(
                slide, spec.get("chart_type", "column"), spec["categories"],
                spec["series"], y=y, h=h,
            )
        if spec.get("source"):
            kit.add_text(
                slide, kit.CONTENT_X, kit.CONTENT_Y + kit.CONTENT_H - 0.05,
                kit.CONTENT_W, 0.25, spec["source"], size=8, color=kit.FAINT,
            )

    elif stype == "stats":
        items = spec.get("items", [])
        if not 2 <= len(items) <= 4:
            warnings.append("stats のタイルは 2〜4 枚が目安。")
        tile_y = y + 0.25 if not lead else y
        kit.add_stats(slide, items[:4], y=tile_y)

    elif stype == "flow":
        steps = spec.get("steps", [])
        if not 3 <= len(steps) <= 5:
            warnings.append("flow のステップは 3〜5 個が目安。")
        kit.add_flow(slide, steps[:5], y=y + (0.2 if not lead else 0))

    elif stype == "comparison":
        kit.add_comparison(
            slide, spec["left"], spec["right"], y=y,
            h=kit.CONTENT_H - (y - kit.CONTENT_Y),
        )

    elif stype == "timeline":
        ms = spec.get("milestones", [])
        if not 3 <= len(ms) <= 6:
            warnings.append("timeline のマイルストーンは 3〜6 個が目安。")
        kit.add_timeline(slide, ms[:6], y=y + 1.2)

    elif stype == "table":
        rows = spec.get("rows", [])
        if len(rows) > 7:
            warnings.append(f"table が {len(rows)} 行。7 行までが目安。")
        kit.add_table(slide, spec["columns"], rows[:8], y=y)

    return slide


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("spec", help="スライド定義 JSON ファイル")
    parser.add_argument("-o", "--output", default="deck.pptx", help="出力 PPTX パス")
    parser.add_argument(
        "--template", default=str(TEMPLATE_PATH), help="テンプレート PPTX パス"
    )
    args = parser.parse_args()

    template_path = Path(args.template).resolve()
    output_path = Path(args.output).resolve()
    if output_path == template_path:
        sys.exit(
            f"出力先がテンプレートと同じパスです: {output_path}\n"
            "テンプレートを破壊するため中止。--output に別のパスを指定してください。"
        )

    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
    slides_spec = spec["slides"]

    valid = set(SLIDE_TYPES) | set(KIT_TYPES)
    unknown = [s["type"] for s in slides_spec if s["type"] not in valid]
    if unknown:
        sys.exit(f"未知のスライドタイプ: {unknown}。有効: {sorted(valid)}")

    prs = Presentation(args.template)
    template_slides = list(prs.slides)
    warnings = []

    for slide_spec in slides_spec:
        stype = slide_spec["type"]
        if stype in SLIDE_TYPES:
            source = template_slides[SLIDE_TYPES[stype]["index"]]
            slide = kit.clone_slide(prs, source)
            build_template_slide(slide, slide_spec, warnings)
        else:
            slide = build_kit_slide(prs, template_slides, slide_spec, warnings)
        if slide_spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = slide_spec["notes"]

    for slide in template_slides:
        kit.delete_slide(prs, slide)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))

    print(f"OK: {len(slides_spec)} 枚のスライドを {output_path} に生成")
    for w in warnings:
        print(f"WARNING: {w}", file=sys.stderr)


if __name__ == "__main__":
    main()
