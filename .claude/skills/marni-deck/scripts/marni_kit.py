"""Marni デザインシステム部品ライブラリ。

テンプレートから抽出したデザイントークン(色・フォント・グリッド)と、
そのトークンに従ってスライド上に組み立てる部品(テキスト・色ブロック・
チャート・KPI タイル・プロセス図・比較・タイムライン・表・全面画像)を提供する。

原則: トークン(色/フォント/ヘッダー様式/フッター)は不変。
レイアウト(部品の組み合わせ方)は内容に合わせて自由。
"""

import copy

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LEGEND_POSITION,
    XL_TICK_LABEL_POSITION,
)
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---------------------------------------------------------------- トークン

FONT = "Helvetica Neue"

RED = RGBColor(218, 18, 18)
BLACK = RGBColor(0, 0, 0)
CAMEL = RGBColor(157, 83, 48)
LIGHT_GREY = RGBColor(210, 210, 210)
DARK_BROWN = RGBColor(64, 32, 35)
LILAC = RGBColor(183, 172, 214)
YELLOW = RGBColor(251, 245, 155)
WHITE = RGBColor(255, 255, 255)
INK = BLACK  # 本文色
FAINT = RGBColor(120, 120, 120)  # 補助情報 (出典など)

# 塗りブロックの上に置く文字色 (fill → text color)
TEXT_ON = {
    RED: WHITE,
    BLACK: WHITE,
    CAMEL: WHITE,
    DARK_BROWN: WHITE,
    LILAC: BLACK,
    YELLOW: BLACK,
    LIGHT_GREY: BLACK,
    WHITE: BLACK,
}

# チャート系列 / タイルの色順
SERIES_COLORS = [RED, BLACK, CAMEL, LILAC, DARK_BROWN, YELLOW]
TILE_COLORS = [RED, BLACK, CAMEL, DARK_BROWN, LILAC]

# グリッド (インチ)。テンプレートの実測値。
SLIDE_W = 10.0
SLIDE_H = 5.625
CONTENT_X = 1.18
CONTENT_Y = 1.18
CONTENT_W = 7.64
CONTENT_H = 3.35  # 下端 4.53
GUTTER = 0.17

# キャンバスのベースとなるテンプレートスライド (0 始まり index)
CANVAS_BASES = {
    "grey": {
        "index": 4,
        "remove": ["Google Shape;166;p22", "Google Shape;167;p22"],
        "header": "Google Shape;165;p22",
        "secnum": "Google Shape;168;p22",
    },
    "red": {
        "index": 2,
        "remove": [],
        "header": "Google Shape;149;p20",
        "secnum": "Google Shape;151;p20",
    },
    "black": {
        "index": 3,
        "remove": [],
        "header": "Google Shape;157;p21",
        "secnum": "Google Shape;159;p21",
    },
}


# ---------------------------------------------------------- 基本プリミティブ


def clone_slide(prs, source_slide):
    """テンプレートスライドを末尾に複製する。

    背景・ロゴ・フッターはレイアウト/マスター側にあり自動で引き継がれる。
    テンプレートのスライド自体は画像 rel を持たない (検証済み)。
    """
    new_slide = prs.slides.add_slide(source_slide.slide_layout)
    for shape in list(new_slide.shapes):
        shape._element.getparent().remove(shape._element)
    for shape in source_slide.shapes:
        new_slide.shapes._spTree.append(copy.deepcopy(shape._element))
    return new_slide


def delete_slide(prs, slide):
    """スライドを sldIdLst と rel の両方から取り除く。"""
    for sld_id in prs.slides._sldIdLst:
        if prs.part.related_part(sld_id.get(qn("r:id"))) is slide.part:
            prs.part.drop_rel(sld_id.get(qn("r:id")))
            prs.slides._sldIdLst.remove(sld_id)
            return
    raise ValueError("slide not found in sldIdLst")


def find_shape(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise KeyError(f"shape {name!r} not found on slide")


def remove_shape(slide, name):
    shape = find_shape(slide, name)
    shape._element.getparent().remove(shape._element)


def fill_text(shape, text):
    """既存シェイプのテキストを書式維持で置換。\n は段落区切り。"""
    tf = shape.text_frame
    tx_body = tf.paragraphs[0]._p.getparent()
    proto = copy.deepcopy(tf.paragraphs[0]._p)
    for extra in proto.findall(qn("a:r"))[1:]:
        proto.remove(extra)
    for p in list(tf.paragraphs):
        p._p.getparent().remove(p._p)
    for line in text.split("\n") if text else [""]:
        p = copy.deepcopy(proto)
        runs = p.findall(qn("a:r"))
        if runs:
            t = runs[0].find(qn("a:t"))
            if t is None:
                t = runs[0].makeelement(qn("a:t"), {})
                runs[0].append(t)
            t.text = line
            if not line:
                p.remove(runs[0])
        tx_body.append(p)


def canvas(prs, template_slides, bg="grey", secnum=None, title=None):
    """ヘッダー付きの空きキャンバスを作る。

    bg="grey" (通常の内容ページ) / "red" / "black" (アクセントページ)。
    title=None ならヘッダーごと取り除く (全面ビジュアル用)。
    """
    base = CANVAS_BASES[bg]
    slide = clone_slide(prs, template_slides[base["index"]])
    for name in base["remove"]:
        remove_shape(slide, name)
    if title is None:
        remove_shape(slide, base["header"])
        remove_shape(slide, base["secnum"])
    else:
        fill_text(find_shape(slide, base["header"]), title)
        fill_text(find_shape(slide, base["secnum"]), secnum or "")
    return slide


# ------------------------------------------------------------------- 部品


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    size=10.5,
    bold=False,
    color=INK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_spacing=1.2,
    wrap=True,
):
    """トークン準拠のテキストボックス。\n は段落区切り。"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        f = run.font
        f.name = FONT
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
    return box


def add_block(slide, x, y, w, h, fill, line_color=None):
    """フラットな色ブロック (パレットページと同じ意匠)。"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape


def add_marker(slide, x, y, size=0.09, color=RED):
    """箇条書き用の小さな色角 (ブランドの色ブロックのミニチュア)。"""
    return add_block(slide, x, y, size, size, color)


def add_cover_picture(slide, path, x, y, w, h):
    """領域全面に画像をクロップ配置 (cover fit)。"""
    from PIL import Image

    with Image.open(path) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), Inches(w), Inches(h)
    )
    box_ratio, img_ratio = w / h, iw / ih
    if img_ratio > box_ratio:
        crop = (1 - box_ratio / img_ratio) / 2
        pic.crop_left = pic.crop_right = crop
    elif img_ratio < box_ratio:
        crop = (1 - img_ratio / box_ratio) / 2
        pic.crop_top = pic.crop_bottom = crop
    return pic


def add_bullets(slide, items, x=CONTENT_X, y=CONTENT_Y, w=CONTENT_W, h=CONTENT_H,
                columns=None):
    """見出し+説明の箇条書き。items: [{"head":..,"text":..} | str, ...]

    4 項目超は自動で 2 カラム。マーカーは赤角。
    """
    items = [{"text": it} if isinstance(it, str) else it for it in items]
    n = len(items)
    if columns is None:
        columns = 2 if n > 4 else 1
    import math

    rows = math.ceil(n / columns)
    col_w = (w - (columns - 1) * 0.4) / columns
    row_h = h / rows
    for i, it in enumerate(items):
        cx = x + (i // rows) * (col_w + 0.4)
        cy = y + (i % rows) * row_h
        head, text = it.get("head"), it.get("text", "")
        add_marker(slide, cx, cy + 0.05)
        tx = cx + 0.22
        if head:
            add_text(slide, tx, cy, col_w - 0.22, 0.3, head, size=12, bold=True)
            add_text(
                slide, tx, cy + 0.3, col_w - 0.22, row_h - 0.34, text, size=10
            )
        else:
            add_text(slide, tx, cy, col_w - 0.22, row_h - 0.08, text, size=10.5)


CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def add_chart(slide, kind, categories, series, x=CONTENT_X, y=CONTENT_Y,
              w=CONTENT_W, h=CONTENT_H):
    """ブランド配色のネイティブチャート。

    kind: column / bar / line / pie / doughnut
    series: [{"name": ..., "values": [...]}, ...] (pie/doughnut は 1 系列)
    """
    data = CategoryChartData()
    data.categories = categories
    for s in series:
        data.add_series(s["name"], s["values"])
    frame = slide.shapes.add_chart(
        CHART_TYPES[kind], Inches(x), Inches(y), Inches(w), Inches(h), data
    )
    chart = frame.chart
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(9)
    chart.font.color.rgb = INK

    circular = kind in ("pie", "doughnut")
    chart.has_legend = circular or len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(9)

    if circular:
        points = chart.plots[0].series[0].points
        for i, pt in enumerate(points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = SERIES_COLORS[i % len(SERIES_COLORS)]
        dl = chart.plots[0].data_labels
        dl.show_value = True
        dl.font.size = Pt(9)
        dl.font.color.rgb = WHITE
    else:
        for i, ser in enumerate(chart.series):
            color = SERIES_COLORS[i % len(SERIES_COLORS)]
            if kind == "line":
                ser.format.line.color.rgb = color
                ser.format.line.width = Pt(2)
                ser.smooth = False
                try:
                    ser.marker.format.fill.solid()
                    ser.marker.format.fill.fore_color.rgb = color
                    ser.marker.format.line.fill.background()
                except Exception:
                    pass
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = color
        try:
            val_ax = chart.value_axis
            val_ax.format.line.fill.background()
            gl = val_ax.major_gridlines.format.line
            gl.color.rgb = LIGHT_GREY
            gl.width = Pt(0.75)
            val_ax.tick_labels.font.size = Pt(9)
            cat_ax = chart.category_axis
            cat_ax.format.line.color.rgb = INK
            cat_ax.tick_labels.font.size = Pt(9)
            # 負の値があってもラベルがバーに重ならないよう軸の外側へ
            if any(v is not None and v < 0 for s in series for v in s["values"]):
                cat_ax.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        except Exception:
            pass
        if kind in ("column", "bar"):
            chart.plots[0].gap_width = 80
    return chart


def add_stats(slide, items, x=CONTENT_X, y=CONTENT_Y, w=CONTENT_W, h=None):
    """KPI タイル 2〜4 枚。items: [{"value","label","desc"?}, ...]"""
    n = len(items)
    tile_h = h or min(1.9, CONTENT_H)
    tile_w = (w - (n - 1) * GUTTER) / n
    for i, it in enumerate(items):
        fill = TILE_COLORS[i % len(TILE_COLORS)]
        tcol = TEXT_ON[fill]
        tx = x + i * (tile_w + GUTTER)
        add_block(slide, tx, y, tile_w, tile_h, fill)
        pad = 0.18
        # 値はタイル幅に収まるフォントサイズへ縮小 (折り返し禁止)
        val = it["value"]
        val_size = 28 if n <= 3 else 24
        est_w = len(val) * val_size * 0.0105  # 太字 Helvetica の概算幅 (in/字/pt)
        max_w = tile_w - 2 * pad
        if est_w > max_w:
            val_size = max(14, int(val_size * max_w / est_w))
        add_text(
            slide, tx + pad, y + pad, max_w, 0.65,
            val, size=val_size, bold=True, color=tcol, wrap=False,
        )
        add_text(
            slide, tx + pad, y + pad + 0.7, tile_w - 2 * pad, 0.32,
            it.get("label", ""), size=10, bold=True, color=tcol,
        )
        if it.get("desc"):
            add_text(
                slide, tx + pad, y + pad + 1.05, tile_w - 2 * pad,
                tile_h - pad - 1.15, it["desc"], size=8.5, color=tcol,
            )


def add_flow(slide, steps, x=CONTENT_X, y=CONTENT_Y, w=CONTENT_W, h=None):
    """番号付きプロセス図 (色ブロック列)。steps: [{"head","text"?}, ...]"""
    n = len(steps)
    block_h = h or 2.4
    block_w = (w - (n - 1) * GUTTER) / n
    for i, st in enumerate(steps):
        fill = TILE_COLORS[i % len(TILE_COLORS)]
        tcol = TEXT_ON[fill]
        bx = x + i * (block_w + GUTTER)
        add_block(slide, bx, y, block_w, block_h, fill)
        pad = 0.15
        add_text(
            slide, bx + pad, y + pad, block_w - 2 * pad, 0.5,
            f"{i + 1:02d}", size=20, bold=True, color=tcol,
        )
        add_text(
            slide, bx + pad, y + pad + 0.55, block_w - 2 * pad, 0.55,
            st.get("head", ""), size=11, bold=True, color=tcol,
        )
        if st.get("text"):
            add_text(
                slide, bx + pad, y + pad + 1.1, block_w - 2 * pad,
                block_h - pad - 1.2, st["text"], size=8.5, color=tcol,
            )


def add_comparison(slide, left, right, x=CONTENT_X, y=CONTENT_Y, w=CONTENT_W,
                   h=CONTENT_H):
    """2 案比較。left/right: {"head", "items": [...], "color"?}"""
    col_w = (w - GUTTER) / 2
    head_h = 0.5
    colors = {"black": BLACK, "camel": CAMEL, "red": RED, "brown": DARK_BROWN,
              "lilac": LILAC}
    for i, side in enumerate((left, right)):
        fill = colors.get(side.get("color", ""), BLACK if i == 0 else CAMEL)
        cx = x + i * (col_w + GUTTER)
        add_block(slide, cx, y, col_w, head_h, fill)
        add_text(
            slide, cx, y + 0.09, col_w, head_h - 0.18, side["head"],
            size=13, bold=True, color=TEXT_ON[fill], align=PP_ALIGN.CENTER,
        )
        body_y = y + head_h + 0.08
        add_block(slide, cx, body_y, col_w, h - head_h - 0.08, WHITE)
        iy = body_y + 0.18
        for item in side.get("items", []):
            add_marker(slide, cx + 0.18, iy + 0.04, color=fill)
            add_text(slide, cx + 0.4, iy, col_w - 0.58, 0.45, item, size=10)
            iy += 0.48


def add_timeline(slide, milestones, x=CONTENT_X, y=None, w=CONTENT_W):
    """横タイムライン。milestones: [{"when","what"}, ...] (3〜6 個)"""
    n = len(milestones)
    line_y = y or (CONTENT_Y + 1.3)
    add_block(slide, x, line_y, w, 0.02, BLACK)
    seg = w / n
    size = 10 if n <= 4 else 9
    for i, ms in enumerate(milestones):
        cx = x + seg * i + seg / 2
        add_block(slide, cx - 0.07, line_y - 0.06, 0.14, 0.14, RED)
        add_text(
            slide, cx - seg / 2 + 0.1, line_y - 0.55, seg - 0.2, 0.35,
            ms["when"], size=size + 1, bold=True, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.BOTTOM,
        )
        add_text(
            slide, cx - seg / 2 + 0.1, line_y + 0.22, seg - 0.2, 1.3,
            ms["what"], size=size, align=PP_ALIGN.CENTER,
        )


def add_table(slide, columns, rows, x=CONTENT_X, y=CONTENT_Y, w=CONTENT_W,
              h=None):
    """ブランド配色の表。ヘッダー黒、行は白/薄グレー交互。"""
    n_rows = len(rows) + 1
    n_cols = len(columns)
    table_h = h or min(CONTENT_H, 0.4 * n_rows)
    frame = slide.shapes.add_table(
        n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(table_h)
    )
    table = frame.table
    table.first_row = False
    table.horz_banding = False
    grey = RGBColor(233, 231, 229)
    for c, name in enumerate(columns):
        _style_cell(table.cell(0, c), name, fill=BLACK, color=WHITE, bold=True)
    for r, row in enumerate(rows):
        fill = WHITE if r % 2 == 0 else grey
        for c in range(n_cols):
            val = str(row[c]) if c < len(row) else ""
            _style_cell(table.cell(r + 1, c), val, fill=fill, color=INK,
                        bold=(c == 0))
    return table


def _style_cell(cell, text, fill, color, bold=False):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    f = run.font
    f.name = FONT
    f.size = Pt(9.5)
    f.bold = bold
    f.color.rgb = color


def add_full_image(slide, path, caption=None):
    """全面ビジュアル + 左下キャプション (黒帯に白文字)。"""
    add_cover_picture(slide, path, 0, 0, SLIDE_W, SLIDE_H)
    if caption:
        cap_w = min(4.5, 0.35 + 0.11 * len(caption) * 2)
        add_block(slide, 0, SLIDE_H - 0.75, cap_w, 0.42, BLACK)
        add_text(
            slide, 0.19, SLIDE_H - 0.68, cap_w - 0.3, 0.3, caption,
            size=10, bold=True, color=WHITE,
        )
